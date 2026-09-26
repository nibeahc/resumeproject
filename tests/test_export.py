"""S8-2 DOCX 내보내기 테스트. LLM 을 안 쓰는 순수 변환 로직이라 실제 docx 파일을 만들고
python-docx 로 다시 읽어서 검증한다 (스텁이 필요 없다)."""
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from lib import export  # noqa: E402


def paragraph_texts(doc):
    return [p.text for p in doc.paragraphs]


def has_eastasia_font(run, name='맑은 고딕'):
    rpr = run._element.find(qn('w:rPr'))
    if rpr is None:
        return False
    rfonts = rpr.find(qn('w:rFonts'))
    return rfonts is not None and rfonts.get(qn('w:eastAsia')) == name


class MarkdownToDocxTest(unittest.TestCase):
    def test_headers_become_heading_styles(self):
        doc = export.markdown_to_docx('# 제목1\n\n## 제목2\n\n### 제목3\n')
        styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(styles, ['Heading 1', 'Heading 2', 'Heading 3'])

    def test_plain_paragraph_preserved(self):
        doc = export.markdown_to_docx('그냥 평범한 문단입니다.')
        self.assertIn('그냥 평범한 문단입니다.', paragraph_texts(doc))

    def test_bold_marker_becomes_bold_run(self):
        doc = export.markdown_to_docx('앞부분 **굵은 부분** 뒷부분')
        p = [p for p in doc.paragraphs if '굵은 부분' in p.text][0]
        bold_runs = [r for r in p.runs if r.text == '굵은 부분']
        self.assertEqual(len(bold_runs), 1)
        self.assertTrue(bold_runs[0].bold)

    def test_horizontal_rule_does_not_crash_and_adds_separator(self):
        doc = export.markdown_to_docx('위\n\n---\n\n아래')
        texts = paragraph_texts(doc)
        self.assertIn('위', texts)
        self.assertIn('아래', texts)
        self.assertTrue(any('─' in t for t in texts))

    def test_bullet_list_items_preserved(self):
        doc = export.markdown_to_docx('- 항목1\n- 항목2\n')
        texts = paragraph_texts(doc)
        self.assertIn('항목1', texts)
        self.assertIn('항목2', texts)

    def test_table_parsed_into_docx_table(self):
        md = '| 번호 | 문항 | 글자수 |\n|---|---|---|\n| 1 | 지원동기 | 500 |\n'
        doc = export.markdown_to_docx(md)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(table.rows[0].cells[1].text, '문항')
        self.assertEqual(table.rows[1].cells[2].text, '500')

    # ── 완료 조건 핵심: 한글 폰트(동아시아 글꼴)가 실제로 지정됐는가 ──
    def test_korean_text_gets_eastasia_font_attribute(self):
        doc = export.markdown_to_docx('# 김채빈 이력서\n\n한글 본문입니다.')
        checked = 0
        for p in doc.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    self.assertTrue(has_eastasia_font(r), f'동아시아 폰트 미지정: {r.text!r}')
                    checked += 1
        self.assertGreater(checked, 0)

    # ── 완료 조건: 문항 제목과 글자수 표기가 유지되는가 (실제 자소서.md 형태로 검증) ──
    def test_coverletter_question_title_and_char_count_preserved(self):
        md = (
            '# 테스트기업 백엔드\n\n'
            '## 지원 동기를 서술하시오\n'
            '글자수: 483자 (공백 포함, 목표 500자)\n\n'
            '실제 자소서 본문입니다.\n\n'
            '---\n'
            '참조한 경험: 백엔드 프로젝트 (2022-백엔드.md)\n'
            '작성일: 2026.09.26\n'
        )
        doc = export.markdown_to_docx(md)
        texts = paragraph_texts(doc)
        self.assertIn('지원 동기를 서술하시오', texts)
        self.assertTrue(any('483자' in t and '500자' in t for t in texts))
        self.assertIn('실제 자소서 본문입니다.', texts)
        self.assertTrue(any('참조한 경험' in t for t in texts))


class ExportDocxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_writes_valid_docx_file_reloadable_by_python_docx(self):
        out = self.tmp / '기업_직무' / '자소서.docx'
        result_path = export.export_docx('# 제목\n\n본문 내용', out)
        self.assertEqual(result_path, out)
        self.assertTrue(out.exists())

        reloaded = Document(str(out))  # 파일 자체가 깨지지 않고 다시 열리는지 확인
        self.assertIn('제목', paragraph_texts(reloaded))
        self.assertIn('본문 내용', paragraph_texts(reloaded))

    def test_creates_parent_directories(self):
        out = self.tmp / 'a' / 'b' / 'c' / 'x.docx'
        export.export_docx('내용', out)
        self.assertTrue(out.exists())


if __name__ == '__main__':
    unittest.main()
